"""Export Service for SuperSearch.

Provides export of indexed documents and OCR results to clean Markdown (.md)
and Microsoft Word (.docx) formats using python-docx.
"""

import os
import re
import html
import datetime

try:
    import docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml, OxmlElement
    from docx.oxml.ns import nsdecls, qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def export_to_markdown(doc: dict, output_path: str) -> str:
    """Exports document content to a standardized Markdown file with YAML frontmatter."""
    title = doc.get("title") or "Tài liệu không tiêu đề"
    orig_path = doc.get("original_path") or doc.get("absolute_original_path") or ""
    doc_type = doc.get("doc_type") or "Tài liệu"
    year = doc.get("year") or "N/A"
    content = doc.get("content") or ""
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    frontmatter = [
        "---",
        f'title: "{title}"',
        f'original_path: "{orig_path}"',
        f'doc_type: "{doc_type}"',
        f'year: "{year}"',
        f'exported_at: "{now_str}"',
        f'exported_by: "SuperSearch"',
        "---",
        "",
        f"# {title}",
        "",
        content.strip(),
        "",
    ]
    output_text = "\n".join(frontmatter)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(output_text)

    return output_path


LATEX_SYMBOLS = {
    r'\alpha': 'α', r'\beta': 'β', r'\gamma': 'γ', r'\delta': 'δ', r'\epsilon': 'ε',
    r'\zeta': 'ζ', r'\eta': 'η', r'\theta': 'θ', r'\iota': 'ι', r'\kappa': 'κ',
    r'\lambda': 'λ', r'\mu': 'μ', r'\nu': 'ν', r'\xi': 'ξ', r'\pi': 'π',
    r'\rho': 'ρ', r'\sigma': 'σ', r'\tau': 'τ', r'\upsilon': 'υ', r'\phi': 'φ',
    r'\chi': 'χ', r'\psi': 'ψ', r'\omega': 'ω',
    r'\Gamma': 'Γ', r'\Delta': 'Δ', r'\Theta': 'Θ', r'\Lambda': 'Λ', r'\Xi': 'Ξ',
    r'\Pi': 'Π', r'\Sigma': 'Σ', r'\Upsilon': 'Υ', r'\Phi': 'Φ', r'\Psi': 'Ψ', r'\Omega': 'Ω',
    r'\le': '≤', r'\leq': '≤', r'\ge': '≥', r'\geq': '≥', r'\ne': '≠', r'\neq': '≠',
    r'\approx': '≈', r'\pm': '±', r'\mp': '∓', r'\times': '×', r'\div': '÷', r'\cdot': '·',
    r'\infty': '∞', r'\partial': '∂', r'\nabla': '∇', r'\in': '∈', r'\notin': '∉',
    r'\forall': '∀', r'\exists': '∃', r'\to': '→', r'\rightarrow': '→', r'\leftarrow': '←',
    r'\Rightarrow': '⇒', r'\Leftarrow': '⇐', r'\Leftrightarrow': '⇔',
    r'\sum': '∑', r'\int': '∫', r'\prod': '∏',
}

MATH_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


def convert_latex_to_omml(latex_code: str, is_block: bool = False) -> str:
    """Converts common LaTeX formulas into native Microsoft Word OMML XML."""
    s = (latex_code or "").strip()
    if s.startswith(r"\(") and s.endswith(r"\)"):
        s = s[2:-2].strip()
    elif s.startswith(r"\[") and s.endswith(r"\]"):
        s = s[2:-2].strip()
        is_block = True
    elif s.startswith("$$") and s.endswith("$$"):
        s = s[2:-2].strip()
        is_block = True
    elif s.startswith("$") and s.endswith("$"):
        s = s[1:-1].strip()

    # Replace standard LaTeX symbol macros
    for sym, val in LATEX_SYMBOLS.items():
        s = s.replace(sym, val)

    # Clean redundant LaTeX markers
    s = re.sub(r'\\left\s*([(\[{|])', r'\1', s)
    s = re.sub(r'\\right\s*([)\]}|])', r'\1', s)
    s = re.sub(r'\\text\{([^{}]+)\}', r'\1', s)
    s = re.sub(r'\\mathrm\{([^{}]+)\}', r'\1', s)
    s = re.sub(r'\\mathbf\{([^{}]+)\}', r'\1', s)

    tokens = []

    def save_xml(m_xml):
        idx = len(tokens)
        tokens.append(m_xml)
        return f"§§MATH{idx}§§"

    def sup_repl(match):
        base = match.group(1).strip()
        sup = match.group(2).strip()
        return save_xml(
            f'<m:sSup><m:e><m:r><m:t>{html.escape(base)}</m:t></m:r></m:e>'
            f'<m:sup><m:r><m:t>{html.escape(sup)}</m:t></m:r></m:sup></m:sSup>'
        )

    def sub_repl(match):
        base = match.group(1).strip()
        sub = match.group(2).strip()
        return save_xml(
            f'<m:sSub><m:e><m:r><m:t>{html.escape(base)}</m:t></m:r></m:e>'
            f'<m:sub><m:r><m:t>{html.escape(sub)}</m:t></m:r></m:sub></m:sSub>'
        )

    def process_sub_sup(text):
        text = re.sub(r'([A-Za-z0-9α-ωΑ-Ω\)\],]+)\^\{([^{}]+)\}', sup_repl, text)
        text = re.sub(r'([A-Za-z0-9α-ωΑ-Ω\)\],]+)\^([A-Za-z0-9α-ωΑ-Ω])', sup_repl, text)
        text = re.sub(r'([A-Za-z0-9α-ωΑ-Ω\)\],]+)_\{([^{}]+)\}', sub_repl, text)
        text = re.sub(r'([A-Za-z0-9α-ωΑ-Ω\)\],]+)_([A-Za-z0-9α-ωΑ-Ω])', sub_repl, text)
        return text

    def segs_to_xml(text):
        segments = re.split(r'(§§MATH\d+§§)', text)
        xml_parts = []
        for seg in segments:
            if not seg:
                continue
            token_match = re.match(r'^§§MATH(\d+)§§$', seg)
            if token_match:
                idx = int(token_match.group(1))
                xml_parts.append(tokens[idx])
            else:
                clean_seg = html.escape(seg)
                xml_parts.append(f'<m:r><m:t>{clean_seg}</m:t></m:r>')
        return "".join(xml_parts)

    # Fractions: \frac{a}{b}
    def frac_repl(match):
        num = process_sub_sup(match.group(1).strip())
        den = process_sub_sup(match.group(2).strip())
        return save_xml(
            f'<m:f><m:num>{segs_to_xml(num)}</m:num>'
            f'<m:den>{segs_to_xml(den)}</m:den></m:f>'
        )

    for _ in range(3):
        if r'\frac' in s:
            s = re.sub(r'\\frac\{([^{}]+)\}\{([^{}]+)\}', frac_repl, s)

    # Radicals: \sqrt[n]{x} or \sqrt{x}
    def rad_deg_repl(match):
        deg = process_sub_sup(match.group(1).strip())
        expr = process_sub_sup(match.group(2).strip())
        return save_xml(
            f'<m:rad><m:deg>{segs_to_xml(deg)}</m:deg>'
            f'<m:e>{segs_to_xml(expr)}</m:e></m:rad>'
        )
    s = re.sub(r'\\sqrt\[([^\]]+)\]\{([^{}]+)\}', rad_deg_repl, s)

    def rad_repl(match):
        expr = process_sub_sup(match.group(1).strip())
        return save_xml(
            f'<m:rad><m:radPr><m:degHide m:val="on"/></m:radPr><m:deg/>'
            f'<m:e>{segs_to_xml(expr)}</m:e></m:rad>'
        )
    s = re.sub(r'\\sqrt\{([^{}]+)\}', rad_repl, s)

    # Process remaining superscripts and subscripts in main string
    s = process_sub_sup(s)

    inner_xml = segs_to_xml(s)
    if is_block:
        return f'<m:oMathPara {MATH_NS}><m:oMath>{inner_xml}</m:oMath></m:oMathPara>'
    return f'<m:oMath {MATH_NS}>{inner_xml}</m:oMath>'


def _add_inline_runs(paragraph, text):
    """Parses inline bold/italic/code markdown and LaTeX math, adding formatted runs and OMML equations."""
    pattern = re.compile(
        r"(\*\*([^*]+)\*\*|\*([^*]+)\*|`([^`]+)`|\\\((.*?)\\\)|(?<!\\)\$(?!\$)([^$\n]+?)(?<!\\)\$)"
    )
    last_idx = 0
    for match in pattern.finditer(text):
        start, end = match.span()
        if start > last_idx:
            paragraph.add_run(text[last_idx:start])

        full_match = match.group(1)
        if full_match.startswith("**") and full_match.endswith("**"):
            run = paragraph.add_run(match.group(2))
            run.bold = True
        elif full_match.startswith("*") and full_match.endswith("*"):
            run = paragraph.add_run(match.group(3))
            run.italic = True
        elif full_match.startswith("`") and full_match.endswith("`"):
            run = paragraph.add_run(match.group(4))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif match.group(5) is not None or match.group(6) is not None:
            math_expr = match.group(5) if match.group(5) is not None else match.group(6)
            inserted = False
            if DOCX_AVAILABLE and parse_xml is not None:
                try:
                    omml_xml = convert_latex_to_omml(math_expr, is_block=False)
                    element = parse_xml(omml_xml)
                    paragraph._p.append(element)
                    inserted = True
                except Exception:
                    inserted = False
            if not inserted:
                run = paragraph.add_run(math_expr)
                run.font.name = "Cambria Math"
                run.italic = True

        last_idx = end

    if last_idx < len(text):
        paragraph.add_run(text[last_idx:])


def export_to_docx(doc: dict, output_path: str) -> str:
    """Exports document content to Microsoft Word (.docx) document."""
    if not DOCX_AVAILABLE:
        raise RuntimeError("Thư viện python-docx chưa được cài đặt.")

    title = doc.get("title") or "Tài liệu không tiêu đề"
    orig_path = doc.get("original_path") or doc.get("absolute_original_path") or ""
    doc_type = doc.get("doc_type") or "Tài liệu"
    year = doc.get("year") or "N/A"
    content = doc.get("content") or ""

    document = docx.Document()

    # Configure Normal Style
    normal_style = document.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Document Header Title
    title_p = document.add_paragraph()
    title_run = title_p.add_run(title)
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(20)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    # Metadata block
    meta_p = document.add_paragraph()
    meta_run = meta_p.add_run(f"Đường dẫn: {orig_path} | Loại: {doc_type} | Năm: {year}")
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    meta_run.italic = True

    # Divider line
    document.add_paragraph("_" * 60)

    lines = content.splitlines()
    i = 0
    n = len(lines)

    while i < n:
        raw_line = lines[i]
        line = raw_line.strip()

        if not line:
            i += 1
            continue

        # Headings
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            heading_text = line.lstrip("#").strip()
            level = min(max(1, level), 3)
            document.add_heading(heading_text, level=level)
            i += 1
            continue
        # Block math \[ ... \] or $$ ... $$
        if line.startswith(r"\[") or line.startswith("$$"):
            block_math_lines = [line]
            is_bracket = line.startswith(r"\[")
            end_token = r"\]" if is_bracket else "$$"
            if not (line.endswith(end_token) and len(line) > 2):
                i += 1
                while i < n and not lines[i].strip().endswith(end_token):
                    block_math_lines.append(lines[i])
                    i += 1
                if i < n:
                    block_math_lines.append(lines[i])
                    i += 1
            else:
                i += 1
            full_math = "\n".join(block_math_lines).strip()
            if full_math.startswith(r"\[") and full_math.endswith(r"\]"):
                clean_expr = full_math[2:-2].strip()
            elif full_math.startswith("$$") and full_math.endswith("$$"):
                clean_expr = full_math[2:-2].strip()
            else:
                clean_expr = full_math

            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inserted = False
            if DOCX_AVAILABLE and parse_xml is not None:
                try:
                    omml_xml = convert_latex_to_omml(clean_expr, is_block=True)
                    p._p.append(parse_xml(omml_xml))
                    inserted = True
                except Exception:
                    inserted = False
            if not inserted:
                r = p.add_run(clean_expr)
                r.font.name = "Cambria Math"
                r.italic = True
            continue

        # Code block ```
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < n:
                i += 1
            code_p = document.add_paragraph()
            code_run = code_p.add_run("\n".join(code_lines))
            code_run.font.name = "Consolas"
            code_run.font.size = Pt(9.5)
            continue

        # Markdown Table
        if "|" in line and i + 1 < n and re.match(r"^\s*\|?[\s\-:|]+\|?\s*$", lines[i + 1]):
            table_lines = [line]
            i += 2  # skip separator line
            while i < n and "|" in lines[i]:
                table_lines.append(lines[i])
                i += 1

            rows_data = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip().strip("|").split("|")]
                rows_data.append(cells)

            if rows_data:
                col_count = max(len(r) for r in rows_data)
                table = document.add_table(rows=len(rows_data), cols=col_count)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = 'Table Grid'
                for row_idx, rdata in enumerate(rows_data):
                    row = table.rows[row_idx]
                    for col_idx in range(col_count):
                        cell_val = rdata[col_idx] if col_idx < len(rdata) else ""
                        cell = row.cells[col_idx]
                        cp = cell.paragraphs[0]
                        crun = cp.add_run(cell_val)
                        if row_idx == 0:
                            crun.bold = True
                document.add_paragraph()  # spacing after table
            continue

        # Bullet list
        if re.match(r"^[\*\-\+]\s+", line):
            text = re.sub(r"^[\*\-\+]\s+", "", line)
            p = document.add_paragraph(style='List Bullet')
            _add_inline_runs(p, text)
            i += 1
            continue

        # Numbered list
        if re.match(r"^\d+\.\s+", line):
            text = re.sub(r"^\d+\.\s+", "", line)
            p = document.add_paragraph(style='List Number')
            _add_inline_runs(p, text)
            i += 1
            continue

        # Blockquote
        if line.startswith(">"):
            quote_text = line.lstrip(">").strip()
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            q_run = p.add_run(quote_text)
            q_run.italic = True
            q_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
            i += 1
            continue

        # Regular paragraph
        p_lines = [line]
        i += 1
        while i < n and lines[i].strip() and not lines[i].strip().startswith("#") and not lines[i].strip().startswith("```") and not lines[i].strip().startswith(">") and not re.match(r"^[\*\-\+]\s+", lines[i].strip()) and not re.match(r"^\d+\.\s+", lines[i].strip()) and not ("|" in lines[i] and i + 1 < n and re.match(r"^\s*\|?[\s\-:|]+\|?\s*$", lines[i + 1])):
            p_lines.append(lines[i].strip())
            i += 1

        p = document.add_paragraph()
        _add_inline_runs(p, " ".join(p_lines))

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    document.save(output_path)
    return output_path
